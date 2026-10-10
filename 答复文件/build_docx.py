#!/usr/bin/env python3
"""Build filing-ready docx from the markdown drafts."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent


def set_run_font(run, name, size_pt, bold=False):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = r.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def add_text_with_bold(paragraph, text, font, size):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, font, size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, font, size, bold=False)


def configure_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    return doc


def is_table_line(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    data = [split_row(r) for r in rows if not re.match(r"^\s*\|?\s*-+", r)]
    if not data:
        return
    cols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(data):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            text = row[j] if j < len(row) else ""
            add_text_with_bold(p, text, "宋体", 10.5 if i else 10.5)
            for run in p.runs:
                run.bold = i == 0 or run.bold
    doc.add_paragraph()


def render_markdown(doc, md_text, title_size=18):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if is_table_line(stripped):
            block = []
            while i < len(lines) and is_table_line(lines[i].strip()):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.first_line_indent = Cm(0)
            add_text_with_bold(p, stripped[2:], "黑体", title_size)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Cm(0)
            add_text_with_bold(p, stripped[3:], "黑体", 14)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Cm(0)
            add_text_with_bold(p, stripped[4:], "黑体", 12)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        if stripped.startswith(">"):
            while i < len(lines) and lines[i].strip().startswith(">"):
                content = lines[i].strip()[1:].strip()
                i += 1
                if not content:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(3)
                add_text_with_bold(p, content, "宋体", 12)
            continue

        if re.match(r"^[-*] ", stripped):
            text = re.sub(r"^[-*] ", "", stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.37)
            add_text_with_bold(p, "• " + text, "宋体", 12)
            i += 1
            continue

        if re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            add_text_with_bold(p, stripped, "宋体", 12)
            i += 1
            continue

        # merge following non-empty plain lines? keep one paragraph per line
        p = doc.add_paragraph()
        if stripped.startswith("**") and stripped.endswith("**") and stripped.count("**") == 2:
            p.paragraph_format.first_line_indent = Cm(0)
        else:
            p.paragraph_format.first_line_indent = Cm(0.74)
        add_text_with_bold(p, stripped, "宋体", 12)
        i += 1


def header_footer(doc, label):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(label)
    set_run_font(run, "宋体", 9)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("申请号：202210956187.7")
    set_run_font(run, "宋体", 9)


def build(md_name, docx_name, header, title_size=18):
    text = (ROOT / md_name).read_text(encoding="utf-8")
    doc = configure_doc()
    header_footer(doc, header)
    render_markdown(doc, text, title_size=title_size)
    out = ROOT / docx_name
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build(
        "意见陈述书.md",
        "意见陈述书.docx",
        "第一次审查意见答复　申请号202210956187.7",
    )
    build(
        "权利要求书替换页.md",
        "权利要求书替换页.docx",
        "权利要求书替换页　申请号202210956187.7",
        title_size=18,
    )
    build(
        "内部撰写说明-不随案提交.md",
        "内部撰写说明-不随案提交.docx",
        "内部文件，不随案提交",
        title_size=16,
    )
